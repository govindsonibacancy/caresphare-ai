"""Deterministic, non-LLM-dependent unit tests for the RAG ingestion
pipeline's building blocks (cleaning, chunking, extraction, embedding
dimension enforcement), plus integration tests confirming what actually
lands in `document_chunks` after a real upload. See
docs/RAG_INGESTION.md.

The unit tests below never call a live embedding service - they either
test pure functions or monkeypatch httpx.post - so they are deterministic
and do not require Ollama to be running. The integration tests at the
bottom go through the real upload API and DO require Ollama (same
dependency test_document_upload.py already has).
"""

import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.main import app
from app.services.documents.chunking import chunk_pages
from app.services.documents.cleaning import clean_text
from app.services.documents.embedding import (
    EmbeddingDimensionMismatch,
    EmbeddingServiceError,
    OllamaEmbeddingService,
)
from app.services.documents.extraction import (
    DocxExtractor,
    ExtractionError,
    MarkdownExtractor,
    TextExtractor,
    extract_text,
)

client = TestClient(app)

HOSPITAL_ADMIN_EMAIL = "vikram.singh@caresphere-demo.example"


# --- cleaning ---------------------------------------------------------


def test_clean_text_normalizes_whitespace_and_collapses_blank_lines():
    raw = "Heading\r\n\r\n\r\nParagraph   with    extra   spaces.\x0cNext page text.\n\n\n\nAnother paragraph."
    cleaned = clean_text(raw)
    assert "\r" not in cleaned
    assert "   " not in cleaned
    assert "\n\n\n" not in cleaned
    assert "Paragraph with extra spaces." in cleaned
    assert "Another paragraph." in cleaned


def test_clean_text_never_rewrites_already_clean_words():
    """Cleaning normalizes whitespace only - it must never alter, summarize,
    or otherwise reinterpret the actual source words."""
    raw = "The patient's dosage is 5mg twice daily.\n\nFollow-up in two weeks."
    assert clean_text(raw) == raw


def test_clean_text_preserves_heading_prefixes():
    raw = "#    Title With Extra Spaces   \n\nBody text."
    cleaned = clean_text(raw)
    assert cleaned.startswith("# Title With Extra Spaces")


# --- chunking -----------------------------------------------------------


def test_chunk_pages_is_deterministic():
    pages = [("# Section One\n\nParagraph text here.\n\n# Section Two\n\nMore text.", 1)]
    first = chunk_pages(pages, chunk_size=1200, overlap=150)
    second = chunk_pages(pages, chunk_size=1200, overlap=150)
    assert first == second


def test_chunk_pages_attributes_content_to_the_heading_active_when_it_was_written():
    """Regression test for a bug where a chunk's outgoing buffer was
    relabeled with the *next* section's heading instead of the one that
    was actually active while its content was accumulated."""
    text_ = (
        "# Emergency Triage Procedure\n\n"
        "Intro text.\n\n"
        "## Initial Assessment\n\n"
        "Upon arrival, patients must be assessed.\n\n"
        "## Escalation\n\n"
        "Escalate promptly."
    )
    chunks = chunk_pages([(text_, None)], chunk_size=1200, overlap=0)

    assessment_chunk = next(c for c in chunks if c.heading == "Initial Assessment")
    assert "Upon arrival, patients must be assessed." in assessment_chunk.content

    escalation_chunk = next(c for c in chunks if c.heading == "Escalation")
    assert "Escalate promptly." in escalation_chunk.content
    assert "Upon arrival" not in escalation_chunk.content


def test_chunk_pages_preserves_source_ordering_via_index():
    pages = [("# A\n\nFirst.\n\n# B\n\nSecond.\n\n# C\n\nThird.", None)]
    chunks = chunk_pages(pages, chunk_size=1200, overlap=0)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert [c.heading for c in chunks] == ["A", "B", "C"]


def test_chunk_pages_preserves_page_numbers():
    """Small pages that fit in one chunk together are allowed to share a
    chunk (tagged with the page the buffer started on) - page numbers are
    only guaranteed to differ once content is large enough to force a
    chunk boundary between pages, as here."""
    pages = [("Page one text with several distinct words in it.", 1), ("Page two text with several distinct words in it.", 2)]
    chunks = chunk_pages(pages, chunk_size=20, overlap=0)
    assert {c.page_number for c in chunks} == {1, 2}


def test_chunk_pages_splits_long_content_without_exceeding_chunk_size():
    words = " ".join(f"word{i}" for i in range(200))
    chunks = chunk_pages([(words, None)], chunk_size=50, overlap=0)
    assert len(chunks) > 1
    assert all(len(c.content) <= 50 for c in chunks)


def test_chunk_pages_overlap_carries_the_previous_chunks_tail_forward():
    words = " ".join(f"word{i}" for i in range(200))
    chunks = chunk_pages([(words, None)], chunk_size=100, overlap=20)
    assert len(chunks) > 1
    for earlier, later in zip(chunks, chunks[1:]):
        tail = earlier.content[-20:]
        assert later.content.startswith(tail)


def test_chunk_pages_on_empty_input_produces_no_chunks():
    assert chunk_pages([], chunk_size=1200, overlap=150) == []
    assert chunk_pages([("", None)], chunk_size=1200, overlap=150) == []


# --- extraction -----------------------------------------------------------


def test_text_extractor_returns_a_single_unpaginated_page():
    extracted = TextExtractor().extract(b"Hello world.")
    assert len(extracted.pages) == 1
    assert extracted.pages[0].page_number is None
    assert extracted.pages[0].text == "Hello world."


def test_markdown_extractor_preserves_heading_syntax_verbatim():
    content = b"# Title\n\nBody text."
    extracted = MarkdownExtractor().extract(content)
    assert extracted.pages[0].text == "# Title\n\nBody text."


def test_docx_extractor_converts_heading_styles_to_markdown_prefixes():
    from docx import Document as DocxDocument

    docx_doc = DocxDocument()
    docx_doc.add_heading("Main Title", level=1)
    docx_doc.add_paragraph("Intro paragraph.")
    docx_doc.add_heading("Sub Section", level=2)
    docx_doc.add_paragraph("Detail paragraph.")
    buffer = io.BytesIO()
    docx_doc.save(buffer)

    extracted = DocxExtractor().extract(buffer.getvalue())
    text_ = extracted.pages[0].text
    assert "# Main Title" in text_
    assert "## Sub Section" in text_
    assert "Intro paragraph." in text_
    assert "Detail paragraph." in text_
    # Heading order relative to body text is preserved.
    assert text_.index("# Main Title") < text_.index("Intro paragraph.") < text_.index("## Sub Section")


def test_extract_text_raises_extraction_error_for_unregistered_extension():
    with pytest.raises(ExtractionError):
        extract_text("exe", b"whatever")


def test_extract_text_wraps_parser_failures_in_a_safe_message():
    """A corrupt file that gets past validation's magic-byte check (this
    call bypasses validation entirely) must never surface a raw parser
    traceback - only a safe, generic message."""
    with pytest.raises(ExtractionError) as excinfo:
        extract_text("docx", b"not actually a docx file")
    message = str(excinfo.value)
    assert "not actually a docx file" not in message
    assert "Traceback" not in message


# --- embedding dimension enforcement (mocked - no live Ollama needed) -----


def _service() -> OllamaEmbeddingService:
    return OllamaEmbeddingService(base_url="http://fake-ollama", model="all-minilm", expected_dimension=384)


def test_embedding_service_accepts_correctly_dimensioned_vectors(monkeypatch):
    class _Response:
        is_error = False

        def json(self):
            return {"embeddings": [[0.1] * 384, [0.2] * 384]}

    monkeypatch.setattr("httpx.post", lambda *a, **k: _Response())
    result = _service().embed(["first chunk", "second chunk"])
    assert len(result) == 2
    assert all(len(vector) == 384 for vector in result)


def test_embedding_service_rejects_wrong_dimension_vectors(monkeypatch):
    class _Response:
        is_error = False

        def json(self):
            return {"embeddings": [[0.1] * 256]}  # e.g. a different model's output

    monkeypatch.setattr("httpx.post", lambda *a, **k: _Response())
    with pytest.raises(EmbeddingDimensionMismatch):
        _service().embed(["a chunk"])


def test_embedding_service_raises_safe_error_on_provider_error_status(monkeypatch):
    class _Response:
        is_error = True
        status_code = 500

        def json(self):
            return {}

    monkeypatch.setattr("httpx.post", lambda *a, **k: _Response())
    with pytest.raises(EmbeddingServiceError):
        _service().embed(["a chunk"])


def test_embedding_service_raises_safe_error_on_unexpected_response_shape(monkeypatch):
    class _Response:
        is_error = False

        def json(self):
            return {"embeddings": [[0.1] * 384]}  # only one vector for two inputs

    monkeypatch.setattr("httpx.post", lambda *a, **k: _Response())
    with pytest.raises(EmbeddingServiceError):
        _service().embed(["first", "second"])


def test_embedding_service_skips_the_network_call_for_empty_input():
    # httpx.post is deliberately left unpatched - a real call here (against
    # a fake, unreachable base_url) would raise, proving embed([]) never
    # calls out at all.
    assert _service().embed([]) == []


# --- DB persistence integration (real upload, requires live Ollama, same
# dependency test_document_upload.py already has) -------------------------


def _upload(headers, *, filename="doc.md", content=b"Some content.", **form):
    files = {"file": (filename, content, "text/markdown")}
    data = {
        "title": form.pop("title", "Ingestion DB Test"),
        "document_type": form.pop("document_type", "HOSPITAL_POLICY"),
        "sensitivity": form.pop("sensitivity", "PUBLIC"),
        **form,
    }
    return client.post("/api/admin/documents", files=files, data=data, headers=headers)


def test_ingested_chunks_are_persisted_with_384_dimensional_vectors_in_order(
    db_engine, auth_as, auth_headers, cleanup_document
):
    auth_as(HOSPITAL_ADMIN_EMAIL)
    content = (
        b"# Heading One\n\nFirst section paragraph.\n\n"
        b"# Heading Two\n\nSecond section paragraph, with enough distinct "
        b"content to be its own chunk."
    )
    response = _upload(auth_headers, content=content, title="DB persistence test")
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "COMPLETED"
    assert body["chunk_count"] >= 2

    with db_engine.connect() as conn:
        rows = conn.execute(
            text(
                "SELECT chunk_index, vector_dims(embedding) AS dims, content "
                "FROM document_chunks WHERE document_id = :id ORDER BY chunk_index"
            ),
            {"id": body["id"]},
        ).mappings().all()

    assert len(rows) == body["chunk_count"]
    assert [row["chunk_index"] for row in rows] == list(range(len(rows)))
    assert all(row["dims"] == 384 for row in rows)
    assert all(row["content"].strip() for row in rows)

    cleanup_document(body["id"])


def test_wrong_dimension_embedding_marks_document_failed_with_zero_chunks(
    db_engine, auth_as, auth_headers, cleanup_document, monkeypatch
):
    """Even if a future embedding-service implementation forgot its own
    dimension check, pgvector's own column type (vector(384)) rejects a
    wrongly-shaped vector at INSERT time - and the pipeline's failure
    handling must still leave zero chunks behind, never a partial set."""

    class _WrongDimensionService:
        def embed(self, texts):
            return [[0.0] * 10 for _ in texts]

    monkeypatch.setattr(
        "app.services.documents.ingestion_service.get_embedding_service", lambda: _WrongDimensionService()
    )

    auth_as(HOSPITAL_ADMIN_EMAIL)
    response = _upload(auth_headers, title="Wrong dimension test", content=b"Some content for embedding.")
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert body["chunk_count"] == 0
    assert body["processing_error"]
    assert "10" not in body["processing_error"]  # safe generic message, not a leaked internal detail

    with db_engine.connect() as conn:
        chunk_count = conn.execute(
            text("SELECT count(*) FROM document_chunks WHERE document_id = :id"), {"id": body["id"]}
        ).scalar_one()
    assert chunk_count == 0

    cleanup_document(body["id"])


def test_ingestion_failure_never_leaves_document_completed(db_engine, auth_as, auth_headers, cleanup_document, monkeypatch):
    """A hard failure partway through embedding must mark the document
    FAILED, never COMPLETED with whatever chunks happened to be inserted
    before the failure."""

    class _FlakyService:
        def embed(self, texts):
            raise EmbeddingServiceError("simulated provider outage")

    monkeypatch.setattr("app.services.documents.ingestion_service.get_embedding_service", lambda: _FlakyService())

    auth_as(HOSPITAL_ADMIN_EMAIL)
    response = _upload(auth_headers, title="Embedding outage test", content=b"Content that will fail to embed.")
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert body["chunk_count"] == 0

    with db_engine.connect() as conn:
        status = conn.execute(text("SELECT status FROM documents WHERE id = :id"), {"id": body["id"]}).scalar_one()
    assert status == "FAILED"

    cleanup_document(body["id"])
