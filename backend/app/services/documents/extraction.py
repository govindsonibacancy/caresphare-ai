"""Format-specific text extraction. Each extractor returns normalized text
plus whatever source structure is actually available (PDF page numbers,
DOCX heading levels) - see docs/RAG_INGESTION.md, "Text extraction".
"""

import dataclasses
import io
from typing import Protocol


@dataclasses.dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None
    text: str


@dataclasses.dataclass(frozen=True)
class ExtractedDocument:
    pages: list[ExtractedPage]

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text)


class DocumentExtractor(Protocol):
    def extract(self, content: bytes) -> ExtractedDocument: ...


class TextExtractor:
    """.txt - no pagination concept; the whole file is one page."""

    def extract(self, content: bytes) -> ExtractedDocument:
        return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=content.decode("utf-8"))])


class MarkdownExtractor:
    """.md - markdown syntax (headings, etc.) is kept as-is rather than
    stripped, since `#`/`##` is exactly what chunking.py looks for to find
    section boundaries."""

    def extract(self, content: bytes) -> ExtractedDocument:
        return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=content.decode("utf-8"))])


class PdfExtractor:
    """.pdf - one ExtractedPage per PDF page, so page numbers survive into
    document_chunks.page_number."""

    def extract(self, content: bytes) -> ExtractedDocument:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [
            ExtractedPage(page_number=index, text=page.extract_text() or "")
            for index, page in enumerate(reader.pages, start=1)
        ]
        return ExtractedDocument(pages=pages)


class DocxExtractor:
    """.docx has no page concept - instead, heading styles are converted to
    markdown-style `#`/`##`/`###` prefixes so chunking.py's section
    detection works the same way it does for a .md source, rather than
    losing that structure entirely."""

    def extract(self, content: bytes) -> ExtractedDocument:
        from docx import Document as DocxDocument

        document = DocxDocument(io.BytesIO(content))
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            if style_name.startswith("heading 1"):
                lines.append(f"# {text}")
            elif style_name.startswith("heading 2"):
                lines.append(f"## {text}")
            elif style_name.startswith("heading 3"):
                lines.append(f"### {text}")
            else:
                lines.append(text)
        return ExtractedDocument(pages=[ExtractedPage(page_number=None, text="\n\n".join(lines))])


_EXTRACTORS: dict[str, DocumentExtractor] = {
    "pdf": PdfExtractor(),
    "docx": DocxExtractor(),
    "txt": TextExtractor(),
    "md": MarkdownExtractor(),
}


class ExtractionError(Exception):
    """Wraps any underlying parser failure into a safe message - the
    ingestion pipeline catches this and marks the document FAILED with
    `str(exc)`, never the raw parser traceback (see
    app/services/documents/ingestion_service.py)."""


def extract_text(extension: str, content: bytes) -> ExtractedDocument:
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise ExtractionError(f"No extractor registered for '.{extension}'.")
    try:
        return extractor.extract(content)
    except Exception as exc:  # deliberately broad: any parser failure -> a safe FAILED status, not a 500
        raise ExtractionError(f"Could not extract text from this {extension.upper()} file.") from exc
